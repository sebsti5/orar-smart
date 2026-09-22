"""Excel export in the UTM layout.

One sheet per study year. Columns: day, time, then one column per group.
Each (day, slot) takes two rows: the top row is the odd week ("imp"), the
bottom row the even week ("par"). A weekly lesson spans both rows; lessons
shared by adjacent groups (stream lectures) are merged horizontally.
"""

from __future__ import annotations

import io
from collections import defaultdict
from collections.abc import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas import Group, InstitutionSetup, PlacedLesson, SolveResult

DAY_NAMES = ["Luni", "Marți", "Miercuri", "Joi", "Vineri", "Sâmbătă", "Duminică"]
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]
KIND_PREFIX = {"lecture": "c.", "seminar": "sem.", "lab": "lab."}
PARITY_SUFFIX = {"odd": "(imp)", "even": "(par)"}
KIND_FILL = {
    "lecture": PatternFill("solid", fgColor="FFF2CC"),
    "seminar": PatternFill("solid", fgColor="E2EFDA"),
    "lab": PatternFill("solid", fgColor="DDEBF7"),
}
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

HEADER_ROW = 4
FIRST_DATA_ROW = 5
FIRST_GROUP_COL = 3
ROWS_PER_SLOT = 2  # odd week on top, even week below


# ------------------------------------------------------------------ text


def lesson_text(setup: InstitutionSetup, lesson: PlacedLesson) -> str:
    subjects = {s.id: s for s in setup.subjects}
    teachers = {t.id: t for t in setup.teachers}
    rooms = {r.id: r for r in setup.rooms}

    subj = subjects.get(lesson.subject_id)
    title = (subj.name or subj.short) if subj else lesson.subject_id
    head = [KIND_PREFIX[lesson.kind], title]
    if lesson.subgroup is not None:
        head.append(f"sg.{lesson.subgroup}")
    if lesson.parity in PARITY_SUFFIX:
        head.append(PARITY_SUFFIX[lesson.parity])
    lines = [" ".join(head)]

    teacher = teachers.get(lesson.teacher_id)
    if teacher is not None:
        lines.append(f"{teacher.title} {teacher.name}".strip())
    else:
        lines.append(lesson.teacher_id)
    room = rooms.get(lesson.room_id) if lesson.room_id else None
    if room is not None:
        lines.append(room.name)
    return "\n".join(lines)


# ------------------------------------------------------------------ layout helpers


def _ordered_groups(setup: InstitutionSetup, year: int) -> list[Group]:
    abbr = {p.id: p.abbreviation for p in setup.programs}
    groups = [g for g in setup.groups if g.year == year]
    return sorted(groups, key=lambda g: (abbr.get(g.program_id, ""), g.name))


def _slot_row(day: int, slot: int, n_slots: int) -> int:
    return FIRST_DATA_ROW + (day * n_slots + slot) * ROWS_PER_SLOT


def _halves(lesson: PlacedLesson) -> tuple[int, ...]:
    return {"odd": (0,), "even": (1,)}.get(lesson.parity, (0, 1))


def _place(lessons: Iterable[PlacedLesson], col_of: dict[str, int], days: int,
           n_slots: int) -> dict[tuple[int, int, int, int], list[PlacedLesson]]:
    """(day, slot, col, half) -> lessons in that cell."""
    cells: dict[tuple[int, int, int, int], list[PlacedLesson]] = defaultdict(list)
    for lesson in lessons:
        if not (0 <= lesson.day < days and 0 <= lesson.slot < n_slots):
            continue
        for gid in lesson.group_ids:
            col = col_of.get(gid)
            if col is None:
                continue
            for half in _halves(lesson):
                cells[(lesson.day, lesson.slot, col, half)].append(lesson)
    return cells


def _key(lessons: list[PlacedLesson]) -> tuple[str, ...]:
    return tuple(sorted(l.id for l in lessons))


def _runs(cols: list[int], key_of) -> list[tuple[int, int, tuple]]:
    """Group consecutive columns with an identical non-empty key -> (first, last, key)."""
    runs: list[tuple[int, int, tuple]] = []
    for col in cols:
        key = key_of(col)
        if key and runs and runs[-1][2] == key and runs[-1][1] == col - 1:
            runs[-1] = (runs[-1][0], col, key)
        elif key:
            runs.append((col, col, key))
    return runs


def _write_block(ws: Worksheet, r1: int, c1: int, r2: int, c2: int, text: str,
                 fill: PatternFill | None) -> None:
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    ws.cell(row=r1, column=c1, value=text)
    if fill is not None:
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(row=r, column=c).fill = fill


# ------------------------------------------------------------------ sheet


def _write_titles(ws: Worksheet, setup: InstitutionSetup, subtitle: str, last_col: int) -> None:
    g = setup.general
    titles = [
        g.name,
        f"ORARUL ACTIVITĂȚILOR DIDACTICE ÎN ANUL UNIVERSITAR {g.academic_year}, "
        f"SEMESTRUL {g.semester}",
        subtitle,
    ]
    for row, text in enumerate(titles, start=1):
        ws.cell(row=row, column=1, value=text)
        ws.cell(row=row, column=1).font = Font(bold=True, size=14 if row < 3 else 12)
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
        if last_col > 1:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)


def _write_frame(ws: Worksheet, setup: InstitutionSetup, groups: list[Group]) -> None:
    g = setup.general
    n_slots = len(g.slots)
    for col, text in enumerate(["Ziua", "Ora"] + [gr.name for gr in groups], start=1):
        c = ws.cell(row=HEADER_ROW, column=col, value=text)
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
    for day in range(g.days_per_week):
        top = _slot_row(day, 0, n_slots)
        bottom = _slot_row(day, n_slots - 1, n_slots) + ROWS_PER_SLOT - 1
        ws.cell(row=top, column=1, value=DAY_NAMES[day % len(DAY_NAMES)]).font = Font(bold=True)
        ws.merge_cells(start_row=top, start_column=1, end_row=bottom, end_column=1)
        for s, slot in enumerate(g.slots):
            r = _slot_row(day, s, n_slots)
            ws.cell(row=r, column=2, value=f"{slot.start}-{slot.end}")
            ws.merge_cells(start_row=r, start_column=2, end_row=r + 1, end_column=2)


def _write_lessons(ws: Worksheet, setup: InstitutionSetup, groups: list[Group],
                   lessons: list[PlacedLesson]) -> None:
    g = setup.general
    n_slots = len(g.slots)
    col_of = {gr.id: FIRST_GROUP_COL + i for i, gr in enumerate(groups)}
    cols = sorted(col_of.values())
    cells = _place(lessons, col_of, g.days_per_week, n_slots)

    def render(ls: list[PlacedLesson]) -> tuple[str, PatternFill | None]:
        ordered = sorted(ls, key=lambda l: (l.subgroup or 0, l.id))
        return "\n".join(lesson_text(setup, l) for l in ordered), KIND_FILL.get(ordered[0].kind)

    for day in range(g.days_per_week):
        for slot in range(n_slots):
            row = _slot_row(day, slot, n_slots)
            top = {c: cells.get((day, slot, c, 0), []) for c in cols}
            bot = {c: cells.get((day, slot, c, 1), []) for c in cols}
            full = {c: _key(top[c]) == _key(bot[c]) for c in cols}

            for c1, c2, _ in _runs(cols, lambda c: _key(top[c]) if full[c] else ()):
                _write_block(ws, row, c1, row + 1, c2, *render(top[c1]))
            for half, source in ((0, top), (1, bot)):
                for c1, c2, _ in _runs(cols, lambda c: () if full[c] else _key(source[c])):
                    _write_block(ws, row + half, c1, row + half, c2, *render(source[c1]))


def _style_grid(ws: Worksheet, last_row: int, last_col: int, n_groups: int) -> None:
    for row in ws.iter_rows(min_row=HEADER_ROW, max_row=last_row, min_col=1, max_col=last_col):
        for c in row:
            c.border = BORDER
            c.alignment = CENTER
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 13
    for i in range(n_groups):
        ws.column_dimensions[get_column_letter(FIRST_GROUP_COL + i)].width = 26
    for r in range(FIRST_DATA_ROW, last_row + 1):
        ws.row_dimensions[r].height = 36
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=FIRST_GROUP_COL)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _year_sheet(wb: Workbook, setup: InstitutionSetup, year: int,
                lessons: list[PlacedLesson]) -> None:
    label = ROMAN[year - 1] if 1 <= year <= len(ROMAN) else str(year)
    ws = wb.create_sheet(f"Anul {label}")
    groups = _ordered_groups(setup, year)
    last_col = FIRST_GROUP_COL + len(groups) - 1
    n_slots = len(setup.general.slots)
    last_row = FIRST_DATA_ROW + setup.general.days_per_week * n_slots * ROWS_PER_SLOT - 1

    _write_titles(ws, setup, f"Anul de studii {label}", last_col)
    _write_frame(ws, setup, groups)
    _write_lessons(ws, setup, groups, lessons)
    _style_grid(ws, last_row, last_col, len(groups))


# ------------------------------------------------------------------ public API


def build_workbook(setup: InstitutionSetup, result: SolveResult) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    years = sorted({g.year for g in setup.groups})
    if not years:
        ws = wb.create_sheet("Orar")
        _write_titles(ws, setup, "Nu există grupe definite.", 1)
        return wb
    for year in years:
        _year_sheet(wb, setup, year, result.lessons)
    return wb


def export_xlsx(setup: InstitutionSetup, result: SolveResult) -> bytes:
    buf = io.BytesIO()
    build_workbook(setup, result).save(buf)
    return buf.getvalue()

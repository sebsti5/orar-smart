import io

import openpyxl

from api_support import sample_setup  # also puts backend/ on sys.path
from app.export import build_workbook, export_xlsx, lesson_text
from app.schemas import PlacedLesson, Room, SolveResult


def _lesson(id, kind, groups, day, slot, teacher="t1", subject="s_am", parity="all",
            subgroup=None, room="r1"):
    return PlacedLesson(id=id, assignment_id=id.split("#")[0], session_index=0, subject_id=subject,
                        kind=kind, teacher_id=teacher, group_ids=groups, subgroup=subgroup,
                        day=day, slot=slot, parity=parity, room_id=room)


def _setup():
    s = sample_setup()
    s.general.name = "Universitatea Tehnică a Moldovei"
    s.rooms.append(Room(id="r3", name="3-120", capacity=30, kind="lab"))
    return s


def _result():
    return SolveResult(status="optimal", lessons=[
        _lesson("a_lec#0", "lecture", ["g1", "g2"], 0, 0),
        _lesson("a_sem1#0", "seminar", ["g1"], 0, 1, teacher="t2", parity="odd", room="r2"),
        _lesson("a_sem2#0", "seminar", ["g1"], 0, 1, teacher="t2", parity="even", room="r2"),
        _lesson("a_lab#0#s1", "lab", ["g2"], 1, 2, subgroup=1, room="r3"),
        _lesson("a_sem3#0", "seminar", ["g3"], 2, 0, teacher="t2", subject="s_pc", room=None),
    ])


def _find(ws, needle):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and needle in c.value:
                return c
    raise AssertionError(f"{needle!r} not found")


def _merged_range_of(ws, cell):
    for rng in ws.merged_cells.ranges:
        if cell.coordinate in rng:
            return rng
    return None


def test_workbook_opens_and_has_sheet_per_year():
    data = export_xlsx(_setup(), _result())
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Anul I", "Anul II"]
    ws = wb["Anul I"]
    assert ws["A1"].value == "Universitatea Tehnică a Moldovei"
    assert ws["A2"].value == "ORARUL ACTIVITĂȚILOR DIDACTICE ÎN ANUL UNIVERSITAR 2026/2027, SEMESTRUL 1"
    header = [c.value for c in ws[4]]
    assert header[:4] == ["Ziua", "Ora", "TI-241", "TI-242"]


def test_stream_lecture_merged_across_groups():
    wb = openpyxl.load_workbook(io.BytesIO(export_xlsx(_setup(), _result())))
    ws = wb["Anul I"]
    c = _find(ws, "c. Analiza Matematică")
    rng = _merged_range_of(ws, c)
    assert rng is not None
    assert rng.min_col == 3 and rng.max_col == 4  # both group columns
    assert rng.max_row - rng.min_row == 1  # both parity rows of the slot
    assert "Stanciu L." in c.value and "3-611" in c.value
    assert c.fill.fgColor.rgb.endswith("FFF2CC")
    assert c.alignment.wrap_text
    assert c.border.left.style == "thin"


def test_biweekly_split_in_two_rows():
    ws = build_workbook(_setup(), _result())["Anul I"]
    odd = _find(ws, "(imp)")
    even = _find(ws, "(par)")
    assert odd.column == even.column == 3
    assert even.row == odd.row + 1
    assert odd.value.startswith("sem. ")


def test_subgroup_label_and_day_merge():
    ws = build_workbook(_setup(), _result())["Anul I"]
    lab = _find(ws, "lab. ")
    assert "sg.1" in lab.value
    assert lab.column == 4
    assert _merged_range_of(ws, lab).max_row - _merged_range_of(ws, lab).min_row == 1
    day = _find(ws, "Luni")
    rng = _merged_range_of(ws, day)
    assert rng.max_row - rng.min_row + 1 == 2 * 7
    t = _find(ws, "08:00-09:30")
    assert _merged_range_of(ws, t).max_row - _merged_range_of(ws, t).min_row == 1


def test_year_two_sheet_room_missing():
    ws = build_workbook(_setup(), _result())["Anul II"]
    c = _find(ws, "sem. Programarea Calculatoarelor")
    assert c.value.count("\n") == 1  # no room line


def test_two_lessons_same_cell_are_joined():
    res = _result()
    res.lessons.append(_lesson("a_lab#0#s2", "lab", ["g2"], 1, 2, subgroup=2, teacher="t2", room="r3"))
    ws = build_workbook(_setup(), res)["Anul I"]
    c = _find(ws, "sg.1")
    assert "sg.2" in c.value


def test_no_groups_and_out_of_range_lessons():
    s = _setup()
    s.groups = []
    wb = build_workbook(s, _result())
    assert wb.sheetnames == ["Orar"]
    s = _setup()
    res = SolveResult(status="optimal", lessons=[_lesson("x#0", "lecture", ["g1", "zz"], 9, 99)])
    wb = build_workbook(s, res)  # ignored silently
    assert "Anul I" in wb.sheetnames


def test_lesson_text_unknown_refs():
    s = _setup()
    l = _lesson("q#0", "lab", ["g1"], 0, 0, teacher="ghost", subject="ghost", room="ghost")
    txt = lesson_text(s, l)
    assert txt.startswith("lab. ")
